
def getRowcount(path,sheetname):
    workbook=openpyxl.load_workbook(path)
    sheet=workbook.active
    return sheet.max_row
def getColumncount(path,sheetname):
    workbook=openpyxl.load_workbook(path)
    sheet=workbook.active
    return sheet.max_column
def readData(path,sheetname,rownum,columnum):
    workbook=openpyxl.load_workbook(path)
    sheet=workbook.active
    return sheet.cell(row=rownum,column=columnum).value

# for r in range(1,rows+1):
#     for c in range(1,cols+1):
#         print (sheet.cell(row=r,column=c).value)
#

from selenium import webdriver
import openpyxl


path = "C:\\Users\\DHANRAJ\\OneDrive\\Documents\\sample.xlsx"

workbook=openpyxl.load_workbook(path)
sheet=workbook.active
rows=sheet.max_row
cols=sheet.max_column

for r in range(1,rows+1):
    for c in range(1,cols+1):
        print(sheet.cell(row=r,column=c).value)


















